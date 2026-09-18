"""Development only: regenerate document inputs without a legacy parser.

pip install -r tools/requirements-fixtures.txt
python tools/generate_fixtures.py --font /path/to/Arial.ttf
Expected results in testdata/expected.json are reviewed separately; never
regenerate goldens from the implementation being tested.
"""
import argparse
from datetime import datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from pypdf import PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
output = ROOT / 'testdata'


def save(name, data):
    (output / name).write_bytes(data)


def serialized(obj):
    data = BytesIO()
    obj.save(data)
    return data.getvalue()


def fixtures(font):
    doc = Document()
    doc.add_paragraph('Подготовить отчёт 😀')
    run = doc.add_paragraph().add_run('до')
    run.add_tab(); run.add_text('после'); run.add_break(); run.add_text('новая строка')
    table = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f'{r}:{c}'
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = 'Вложенный текст'
    doc.add_paragraph('Абзац после таблицы')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Вторая таблица'
    save('text.docx', serialized(doc))

    book = Workbook()
    sheet = book.active; sheet.title = 'Задачи'
    sheet.append(['Задача', 'Приоритет', None, ''])
    sheet.append(['Позвонить 😀', 'Высокий', True, False, 42, 1.25, -2, '=SUM(E2:F2)', '#DIV/0!'])
    sheet.append([datetime(2026, 9, 17, 15, 30, 12, 123000), time(12, 34, 56, 789000), timedelta(days=2, hours=3, minutes=4), 0.125])
    sheet['D3'].number_format = '0.00%'
    sheet['A4'] = 'Объединено'; sheet.merge_cells('A4:C4')
    sheet['A8'].fill = PatternFill('solid', fgColor='FF0000')
    sheet['A9'] = '=HYPERLINK("https://example.invalid", "Не загружать")'
    book.create_sheet('Пустой')
    save('values.xlsx', serialized(book))
    book.epoch = CALENDAR_MAC_1904
    save('dates1904.xlsx', serialized(book))
    # Shared formulas must be expanded even without cached calculated values.
    source = BytesIO(serialized(book)); replacement = BytesIO()
    with ZipFile(source) as src, ZipFile(replacement, 'w', ZIP_DEFLATED) as dest:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'ref="A1:I9"', b'ref="A1:I11"')
                data = data.replace(b'</sheetData>', b'<row r="10"><c r="A10"><f t="shared" ref="A10:A11" si="0">E2+1</f><v></v></c></row><row r="11"><c r="A11"><f t="shared" si="0"/><v></v></c></row></sheetData>')
            dest.writestr(item, data)
    save('shared-formulas.xlsx', replacement.getvalue())

    writer = PdfWriter(); writer.add_blank_page(width=200, height=200)
    buf = BytesIO(); writer.write(buf); save('blank.pdf', buf.getvalue())
    for password, name in [('secret', 'encrypted.pdf'), ('', 'encrypted-empty.pdf')]:
        secured = PdfWriter(); secured.add_blank_page(width=200, height=200); secured.encrypt(password)
        buf = BytesIO(); secured.write(buf); (output / name).write_bytes(buf.getvalue())
    pdfmetrics.registerFont(TTFont('FixtureFont', font))
    buf = BytesIO(); pdf = canvas.Canvas(buf, invariant=1)
    pdf.setFont('FixtureFont', 12); pdf.drawString(50, 750, 'Задача: подготовить отчёт'); pdf.drawString(50, 725, 'Deadline: Friday')
    pdf.showPage(); pdf.setFont('FixtureFont', 12); pdf.drawString(50, 750, 'Вторая страница'); pdf.save()
    save('text.pdf', buf.getvalue())
    buf = BytesIO(); pdf = canvas.Canvas(buf, invariant=1)
    pdf.beginForm('nested'); pdf.setFont('FixtureFont', 12); pdf.drawString(20, 100, 'Текст внутри формы'); pdf.endForm()
    pdf.doForm('nested'); pdf.save()
    save('form.pdf', buf.getvalue())


if __name__ == '__main__':
    args = argparse.ArgumentParser(); args.add_argument('--font', required=True)
    fixtures(args.parse_args().font)
