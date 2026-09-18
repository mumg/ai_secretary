from io import BytesIO
from unittest import TestCase

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

import json
import os
from pathlib import Path
import subprocess
import tempfile

ROOT = Path(__file__).resolve().parents[1]


class DocumentExtractorTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.directory.cleanup)
        cls.binary = Path(cls.directory.name) / ("document-parser.exe" if os.name == "nt" else "document-parser")
        subprocess.run(["go", "build", "-o", str(cls.binary), "./cmd/document-parser"], cwd=ROOT, check=True)

    def extract(self, suffix, data):
        result = subprocess.run([str(self.binary), "parse", "--suffix", suffix], input=data, capture_output=True, check=True)
        parsed = json.loads(result.stdout)
        return parsed["text"], parsed["metadata"]

    def test_docx_paragraph_and_table(self) -> None:
        document = Document()
        document.add_paragraph("Подготовить отчёт")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Срок"
        table.cell(0, 1).text = "Пятница"
        data = BytesIO()
        document.save(data)

        output, metadata = self.extract(".docx", data.getvalue())
        self.assertIn("Подготовить отчёт", output)
        self.assertIn("Срок\tПятница", output)
        self.assertEqual(metadata["tables"], 1)

    def test_xlsx_reads_values_without_external_links(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Задачи"
        sheet.append(["Задача", "Приоритет"])
        sheet.append(["Позвонить", "Высокий"])
        data = BytesIO()
        workbook.save(data)
        workbook.close()

        output, metadata = self.extract(".xlsx", data.getvalue())
        self.assertIn("[Лист: Задачи]", output)
        self.assertIn("Позвонить\tВысокий", output)
        self.assertEqual(metadata["non_empty_cells"], 4)

    def test_pdf_reports_page_count(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        data = BytesIO()
        writer.write(data)

        output, metadata = self.extract(".pdf", data.getvalue())
        self.assertIn("[Страница 1]", output)
        self.assertEqual(metadata["pages"], 1)
