from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from pypdf import PdfReader

MAX_BYTES = int(os.getenv("MAX_BYTES", str(25 * 1024 * 1024)))
MAX_CHARACTERS = int(os.getenv("MAX_CHARACTERS", "100000"))
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx"}

app = FastAPI(title="Improver document parser", docs_url=None, redoc_url=None)


class ExtractionResult(BaseModel):
    text: str
    media_type: str
    truncated: bool
    metadata: dict[str, int | str]


class BoundedText:
    def __init__(self, limit: int):
        self.limit = limit
        self.parts: list[str] = []
        self.length = 0
        self.truncated = False

    def add(self, value: object) -> bool:
        if value is None or self.truncated:
            return not self.truncated
        text = str(value).strip()
        if not text:
            return True
        available = self.limit - self.length
        if available <= 0:
            self.truncated = True
            return False
        fragment = text[:available]
        self.parts.append(fragment)
        self.length += len(fragment) + 1
        if len(text) > available:
            self.truncated = True
            return False
        return True

    def render(self) -> str:
        return "\n".join(self.parts)


def extract_pdf(data: bytes) -> tuple[BoundedText, dict[str, int | str]]:
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise ValueError("Password-protected PDF is not supported")
    output = BoundedText(MAX_CHARACTERS)
    for number, page in enumerate(reader.pages, start=1):
        if not output.add(f"[Страница {number}]"):
            break
        if not output.add(page.extract_text() or ""):
            break
    return output, {"pages": len(reader.pages)}


def extract_docx(data: bytes) -> tuple[BoundedText, dict[str, int | str]]:
    document = Document(BytesIO(data))
    output = BoundedText(MAX_CHARACTERS)
    for paragraph in document.paragraphs:
        if not output.add(paragraph.text):
            break
    table_count = 0
    for table_count, table in enumerate(document.tables, start=1):
        if not output.add(f"[Таблица {table_count}]"):
            break
        for row in table.rows:
            if not output.add("\t".join(cell.text for cell in row.cells)):
                break
        if output.truncated:
            break
    return output, {"paragraphs": len(document.paragraphs), "tables": table_count}


def extract_xlsx(data: bytes) -> tuple[BoundedText, dict[str, int | str]]:
    workbook = load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False, rich_text=False
    )
    output = BoundedText(MAX_CHARACTERS)
    cell_count = 0
    for sheet in workbook.worksheets:
        if not output.add(f"[Лист: {sheet.title}]"):
            break
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            cell_count += len(values)
            if values and not output.add("\t".join(values)):
                break
        if output.truncated:
            break
    workbook.close()
    return output, {"sheets": len(workbook.sheetnames), "non_empty_cells": cell_count}


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/extract", response_model=ExtractionResult)
async def extract(file: UploadFile = File(...)) -> ExtractionResult:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise HTTPException(status_code=415, detail=f"Unsupported file type: {suffix or 'unknown'}")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="File is too large")
    try:
        if suffix == ".pdf":
            output, metadata = extract_pdf(data)
        elif suffix == ".docx":
            output, metadata = extract_docx(data)
        else:
            output, metadata = extract_xlsx(data)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Unable to parse document: {exc}") from None
    return ExtractionResult(
        text=output.render(),
        media_type=file.content_type or "application/octet-stream",
        truncated=output.truncated,
        metadata=metadata,
    )

